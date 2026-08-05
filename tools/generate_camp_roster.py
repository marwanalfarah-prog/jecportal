# -*- coding: utf-8 -*-
"""
Generate the "camp participants roster" spreadsheet for an event, reproducing the
look of `Elementary Camp Participants 2026.xlsx` (repo root) from live event data.

Standalone tool — NOT part of the running web app and it NEVER writes to the
project's data. It reads data/events.json + the CSV store through the backend's
own export helpers, so names, genders, health conditions and notes match what the
app shows.

Output is one right-to-left `.xlsx` sheet named `Actual`:
  • header block: year motto (+ motto logo & JEC Jordan logo), event title, location & dates
  • participants section: one block per youth group (alphabetical by short name),
    girls in the البنات column, boys in the الشباب column, live IF/SUM/fee formulas
  • الأمانة واللجان section: committee blocks (alphabetical), fee = 0

Two things are resolved AT EXPORT TIME (as agreed):
  1. Gender — taken from the persons store; if any included person has no gender,
     you are forced to pick it before the file is written.
  2. Hull — a committee member listed under several hulls is placed in exactly one;
     you pick which. (Guests are auto-routed by title/reason.)
Choices are saved next to the output as `<out>.resolutions.json` and reused on
re-runs, so you only answer once. Non-interactive runs read `--resolutions FILE`.

Usage (run from anywhere):
    python tools/generate_camp_roster.py --event EVT000002
    python tools/generate_camp_roster.py --event EVT000002 --fee 25 --out C:/path/roster.xlsx
    python tools/generate_camp_roster.py --list
    python tools/generate_camp_roster.py --event EVT000002 --resolutions res.json --no-open
"""
import os
import sys
import json
import argparse
import webbrowser

# ── locate the repo / backend regardless of where we're invoked from ──────────
_THIS_DIR = os.path.dirname(os.path.abspath(__file__))
_REPO     = os.path.abspath(os.path.join(_THIS_DIR, '..'))
_BACKEND  = os.path.join(_REPO, 'backend')
_DATA     = os.path.join(_REPO, 'data')

sys.path.insert(0, _BACKEND)
# CSV helpers resolve some paths against the process cwd — run from backend/.
os.chdir(_BACKEND)

# Arabic output must survive the Windows console's default codepage.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

import core.state as S           # noqa: E402
import core.routes_events as RE  # noqa: E402

import openpyxl                                    # noqa: E402
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side  # noqa: E402
from openpyxl.styles.colors import Color           # noqa: E402
from openpyxl.utils import get_column_letter       # noqa: E402
try:
    from openpyxl.drawing.image import Image as XLImage  # noqa: E402
except Exception:                                        # pragma: no cover
    XLImage = None


# ── constants that mirror the original workbook's look ────────────────────────
BLUE_FILL  = 'FFC6D9F0'   # participants header row
GREEN_FILL = 'FFEAF1DD'   # committees section header row
TITLE_FONT = 'DG Ghayaty'
BODY_FONT  = 'Calibri'
FEMALE, MALE = 'أنثى', 'ذكر'

# Which registration states count as "confirmed & attending".
_ATTENDING = ('registered', 'attended')
_PRIEST_TITLES = ('الأب', 'الأخت', 'المطران', 'الخوري', 'الأنبا', 'القس',
                  'الشماس', 'الأرشمندريت', 'المونسنيور', 'الراهب', 'الراهبة')
# Titles kept when a name is shortened to "first + last".
_NAME_TITLES = _PRIEST_TITLES + ('الدكتور', 'الدكتورة', 'المرشد', 'المرشدة',
                                 'الأستاذ', 'الأستاذة', 'الشيخ')

# Column layout (1=A … 10=J). Widths/labels copied from the original file.
COL = dict(spacer=1, num=2, group=3, total=4, girl=5, gcount=6,
           boy=7, bcount=8, fee=9, notes=10)
COL_WIDTHS = {'A': 1.1, 'B': 3.4, 'C': 16.1, 'D': 10.1, 'E': 22.7,
              'F': 8.9, 'G': 25.5, 'H': 8.9, 'I': 11.2, 'J': 40.1}
HEADERS = ['', '', 'الشبيبة', 'العدد الكامل', 'البنات', 'العدد',
           'الشباب', 'العدد', 'الاشتراك', 'ملاحظات ']

THIN = Side(style='thin')
MEDIUM = Side(style='medium')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

# Salmon bar between sections — theme accent2, lightened (matches the original file).
SEP_FILL = PatternFill('solid', fgColor=Color(theme=5, tint=0.6))

# Header banner palette (a cleaner take on the original three-line header).
BANNER_FILL = 'FFF3F7FC'   # very light blue, sits between the two logos
INK_TITLE   = 'FF1F3864'   # deep blue for the event name (the hero line)
INK_MOTTO   = 'FF9A6B18'   # warm gold for the motto, echoing the motto logo
INK_SUB     = 'FF595959'   # muted grey for location & dates


# ── read-only state load (skip init_state's schema-normalise + re-save) ───────
def _load_state_readonly():
    S.load()
    for fn in ('_refresh_youth_group_indexes', '_load_unreg_store'):
        try:
            getattr(S, fn)()
        except Exception:
            pass


# ── small helpers ─────────────────────────────────────────────────────────────
def _txt(v):
    return ('' if v is None else str(v)).strip()


def _yg_short(label):
    """`شبيبة مار إلياس - الوهادنة` → `الوهادنة` (town after the last ' - ')."""
    label = _txt(label)
    if ' - ' in label:
        return label.rsplit(' - ', 1)[1].strip()
    if label.startswith('شبيبة '):
        return label[len('شبيبة '):].strip()
    return label


def _sort_key(s):
    """Deterministic Arabic-ish alphabetical key (strip diacritics/tatweel)."""
    import unicodedata
    s = _txt(s)
    s = ''.join(c for c in unicodedata.normalize('NFKD', s)
                if not unicodedata.combining(c))
    return s.replace('ـ', '')


# Surname prefixes that belong to the last name (keep them attached when shortening).
_SURNAME_PREFIXES = ('أبو', 'ابو', 'عبد', 'عشم', 'آل', 'ابن', 'بني', 'أم', 'ام', 'بو')


def _short_name_heuristic(name):
    """Shorten a stored full-name string to first + last when the structured
    fields aren't available (unregistered people). Clergy titles and compound
    surnames are preserved: `الأب وجدي معين جمال الطوال` → `الأب وجدي الطوال`,
    `عيسى يزيد عيسى أبو حنك` → `عيسى أبو حنك`."""
    tokens = _txt(name).split()
    if not tokens:
        return ''
    title = tokens.pop(0) if tokens[0] in _NAME_TITLES else ''
    if not tokens:
        return title
    # a compound surname keeps its prefix (e.g. أبو حنك, عشم الله)
    last = tokens[-1]
    if len(tokens) >= 2 and tokens[-2] in _SURNAME_PREFIXES:
        last = f'{tokens[-2]} {tokens[-1]}'
    first = tokens[0]
    parts = [p for p in (title, first, last) if p]
    # avoid duplicating when first and last collapse to the same single token
    if len(parts) >= 2 and parts[-1] == parts[-2]:
        parts = parts[:-1]
    return ' '.join(parts)


def _person_short_name(entry):
    """First + last name for a registration entry. Uses the person's structured
    ar_first_name / ar_last_name (surname is compound-safe) when available,
    falling back to a heuristic on the stored name string for unregistered people."""
    pid = entry.get('person_id')
    person = {}
    try:
        _st, _store, person = RE._resolve_profile_source(
            pid, fallback_unregistered=bool(entry.get('is_unregistered')))
    except Exception:
        person = {}
    title = _txt(person.get('title'))
    first = _txt(person.get('ar_first_name'))
    last  = _txt(person.get('ar_last_name'))
    if first or last:
        # clergy titles are often only on the stored name string, not the record
        if not title:
            toks = _txt(entry.get('name')).split()
            if toks and toks[0] in _NAME_TITLES:
                title = toks[0]
        return ' '.join(p for p in (title, first, last) if p)
    return _short_name_heuristic(_txt(entry.get('name')))


def _combine_notes(reg_notes, health, profile_notes):
    parts = [_txt(reg_notes), _txt(health), _txt(profile_notes)]
    return ' — '.join(p for p in parts if p)


def _is_attending(att):
    return (RE._normalize_attendance_status(att) or RE.ATTENDANCE_DEFAULT_STATUS) in _ATTENDING


def _starts_with_title(name):
    n = _txt(name)
    return any(n.startswith(t + ' ') or n == t for t in _PRIEST_TITLES)


# ── motto / header text sourced from the same config the app uses ─────────────
def _load_bible_abbr():
    path = os.path.join(_DATA, 'bible_books', 'bible_books.json')
    abbr = {}
    try:
        with open(path, encoding='utf-8') as f:
            data = json.load(f)

        def walk(o):
            if isinstance(o, dict):
                if o.get('id') and o.get('abbr'):
                    abbr[o['id']] = o['abbr']
                for v in o.values():
                    walk(v)
            elif isinstance(o, list):
                for v in o:
                    walk(v)
        walk(data)
    except Exception:
        pass
    return abbr


def _active_motto(jec_year):
    """Return {'text': '"..." (1 بط 3: 15)', 'logo': <abs path or None>} for the year."""
    import csv
    path = os.path.join(_DATA, 'JECJordanData', 'scd_mottos.csv')
    best = None
    try:
        with open(path, encoding='utf-8-sig') as f:
            for row in csv.DictReader(f):
                if _txt(row.get('year_label')) != _txt(jec_year):
                    continue
                if _txt(row.get('scd_currently_active_flag')).lower() in ('false', '0', 'no', 'f'):
                    continue
                best = row
                if _txt(row.get('targets_jec_jordan')).lower() in ('true', '1', 'yes'):
                    break
    except Exception:
        pass
    if not best:
        return {'text': '', 'logo': None}

    title = _txt(best.get('title'))
    abbr = _load_bible_abbr().get(_txt(best.get('bible_book_id')), '')
    chap = _txt(best.get('bible_section_start'))
    v1 = _txt(best.get('bible_verse_start'))
    v2 = _txt(best.get('bible_verse_end'))
    ref = ''
    if abbr and chap and v1:
        verses = v1 if (not v2 or v2 == v1) else f'{v1}-{v2}'
        ref = f'({abbr} {chap}: {verses})'
    text = f'"{title}" {ref}'.strip() if title else ''

    logo = None
    fn = _txt(best.get('logo_file_name'))
    if fn:
        cand = os.path.join(_DATA, 'photos', 'logos', 'mottos', fn)
        if os.path.exists(cand):
            logo = cand
    return {'text': text, 'logo': logo}


def _format_date_range(evt):
    from datetime import datetime

    def parse(s):
        s = _txt(s)
        for fmt in ('%Y-%m-%dT%H:%M', '%Y-%m-%dT%H:%M:%S', '%Y-%m-%d'):
            try:
                return datetime.strptime(s, fmt)
            except Exception:
                continue
        return None
    a, b = parse(evt.get('start_datetime')), parse(evt.get('end_datetime'))
    if not a or not b:
        return ''
    if a.year == b.year and a.month == b.month:
        return f'{a.day} - {b.day} / {a.month} / {a.year}'
    if a.year == b.year:
        return f'{a.day} / {a.month} - {b.day} / {b.month} / {a.year}'
    return f'{a.day}/{a.month}/{a.year} - {b.day}/{b.month}/{b.year}'


def _location_name(evt):
    locs = evt.get('locations') or []
    if not locs:
        return ''
    name = _txt(locs[0].get('name'))
    return name.split(' - ', 1)[0].strip() if ' - ' in name else name


def _jec_logo_path():
    p = os.path.join(_DATA, 'photos', 'logos', 'JECJordanLogo.png')
    return p if os.path.exists(p) else None


# ── build the roster model (blocks of people) from the event ──────────────────
def _people_from(rows, entries):
    """Zip enriched export rows back onto their registration entries."""
    out = []
    for row, entry in zip(rows, entries):
        out.append({
            'reg_id': entry.get('id'),
            'name':   _person_short_name(entry),
            'gender': _txt(row.get('الجنس')) or _txt(entry.get('gender')),
            'notes':  _combine_notes(row.get('ملاحظات التسجيل'),
                                     row.get('الحالات الصحية'),
                                     row.get('ملاحظات الملف الشخصي')),
            'entry':  entry,
        })
    return out


def build_model(evt, fee, resolver):
    reg = evt.get('registration', {}) or {}

    mem = _people_from(RE._build_member_export_rows(evt), reg.get('members', []))
    sup = _people_from(RE._build_supervisor_export_rows(evt), reg.get('supervisors', []))
    gsc = _people_from(RE._build_gs_committee_export_rows(evt), reg.get('gs_committee', []))

    # ── participants: keep confirmed & attending, group by YG short name ──────
    part_groups = {}
    for p in mem:
        e = p['entry']
        conf = _txt(e.get('confirmation_status')) or 'confirmed'
        if conf != 'confirmed' or not _is_attending(e.get('attendance_status')):
            continue
        label = _txt(e.get('youth_group_label')) or _txt(RE._lookup_yg_label(e.get('youth_group_id')))
        short = _yg_short(label)
        if not short:
            continue  # members without a youth group have no effect on the output
        p['gender'] = resolver.gender(p)
        part_groups.setdefault(short, []).append(p)

    participants = [{'title': name, 'people': ppl, 'fee': fee}
                    for name, ppl in part_groups.items()]
    participants.sort(key=lambda b: _sort_key(b['title']))

    # ── committees section: gs hulls + supervisors + guests, all fee 0 ────────
    committee = {}

    def add(block, person):
        committee.setdefault(block, []).append(person)

    for p in gsc:
        e = p['entry']
        if _txt(e.get('status')) != RE.FINAL_PERSON_APPROVED_STATUS or not _is_attending(e.get('attendance_status')):
            continue
        hulls = [h for h in (e.get('hulls') or []) if _txt(h)]
        if not hulls:
            continue
        hull = hulls[0] if len(hulls) == 1 else resolver.hull(p, hulls)
        p['gender'] = resolver.gender(p)
        add(hull, p)

    for p in sup:
        e = p['entry']
        if _txt(e.get('status')) != RE.FINAL_PERSON_APPROVED_STATUS or not _is_attending(e.get('attendance_status')):
            continue
        p['gender'] = resolver.gender(p)
        add('مسؤولي الفرق', p)

    for e in reg.get('guests', []):
        if not _is_attending(e.get('attendance_status')):
            continue
        reason = _txt(e.get('reason'))
        name = _txt(e.get('name'))
        if 'الكاريتاس' in reason:
            block = 'الكاريتاس'
        elif _starts_with_title(name):
            block = 'كهنة وشمامسة'
        else:
            block = 'محاضرون'
        person = {'reg_id': e.get('id'), 'name': _person_short_name(e),
                  'gender': _txt(e.get('gender')), 'notes': reason, 'entry': e}
        person['gender'] = resolver.gender(person)
        add(block, person)

    committees = [{'title': name, 'people': ppl, 'fee': 0}
                  for name, ppl in committee.items()]
    committees.sort(key=lambda b: _sort_key(b['title']))

    motto = _active_motto(evt.get('jec_year'))
    return {
        'motto_text': motto['text'],
        'motto_logo': motto['logo'],
        'jec_logo':   _jec_logo_path(),
        'title':      f"{_txt(evt.get('base_title'))} {_txt(evt.get('jec_year'))}".strip(),
        'location':   f"{_location_name(evt)} {_format_date_range(evt)}".strip(),
        'participants': participants,
        'committees':   committees,
    }


# ── resolution of gender + multi-hull (interactive, cached in a sidecar) ───────
class Resolver:
    def __init__(self, saved, interactive):
        self.saved = {'gender': dict(saved.get('gender', {})),
                      'hull':   dict(saved.get('hull', {}))}
        self.interactive = interactive
        self.pending = []          # human-readable list of what could not be resolved

    def gender(self, person):
        g = _txt(person.get('gender'))
        if g in (MALE, FEMALE):
            return g
        rid = person.get('reg_id')
        if rid in self.saved['gender']:
            return self.saved['gender'][rid]
        if self.interactive:
            g = self._ask_gender(person)
            self.saved['gender'][rid] = g
            return g
        self.pending.append(f"gender  · {person.get('name')}  (reg {rid})")
        return ''      # placeholder; run will abort before writing

    def hull(self, person, hulls):
        rid = person.get('reg_id')
        if rid in self.saved['hull'] and self.saved['hull'][rid] in hulls:
            return self.saved['hull'][rid]
        if self.interactive:
            h = self._ask_hull(person, hulls)
            self.saved['hull'][rid] = h
            return h
        self.pending.append(f"hull    · {person.get('name')}  (reg {rid}) → {' / '.join(hulls)}")
        return hulls[0]

    # -- prompts ---------------------------------------------------------------
    def _ask_gender(self, person):
        print(f"\n  ‹؟› لا يوجد جنس محدّد لـ: {person.get('name')}")
        while True:
            ans = input("      اختر  [1] أنثى   [2] ذكر  : ").strip()
            if ans in ('1', 'أنثى', 'انثى', 'f', 'F'):
                return FEMALE
            if ans in ('2', 'ذكر', 'm', 'M'):
                return MALE

    def _ask_hull(self, person, hulls):
        print(f"\n  ‹؟› {person.get('name')} مدرج في أكثر من هيكل — اختر واحدًا:")
        for i, h in enumerate(hulls, 1):
            print(f"      [{i}] {h}")
        while True:
            ans = input("      رقم الاختيار: ").strip()
            if ans.isdigit() and 1 <= int(ans) <= len(hulls):
                return hulls[int(ans) - 1]


# ── write the workbook ─────────────────────────────────────────────────────────
def _header_banner(ws, motto, title, location):
    """Three-line header banner (motto / event title / location) spanning B:J,
    styled as one unit: a light fill between the logos, the event name as the
    hero line, and a medium rule under it to separate from the table."""
    lines = [
        (2, motto,    14, INK_MOTTO, 26),   # motto — warm gold
        (3, title,    20, INK_TITLE, 34),   # event name — hero
        (4, location, 12, INK_SUB,   22),   # location & dates — muted
    ]
    last = lines[-1][0]
    for row, text, size, ink, height in lines:
        ws.merge_cells(start_row=row, start_column=COL['num'],
                       end_row=row, end_column=COL['notes'])
        c = ws.cell(row=row, column=COL['num'], value=text)
        c.font = Font(name=TITLE_FONT, size=size, bold=True, color=ink)
        c.alignment = CENTER
        ws.row_dimensions[row].height = height
        for col in range(COL['num'], COL['notes'] + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = PatternFill('solid', fgColor=BANNER_FILL)
            # outer thin box around the whole banner; medium rule under the last line
            cell.border = Border(
                left=THIN if col == COL['num'] else None,
                right=THIN if col == COL['notes'] else None,
                top=THIN if row == 2 else None,
                bottom=MEDIUM if row == last else None)


def _separator_row(ws, row, height=7.5):
    """Coloured spacer bar between sections (theme accent2, lightened)."""
    ws.row_dimensions[row].height = height
    for col in range(COL['num'], COL['notes'] + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = SEP_FILL
        cell.border = BORDER


def _style_row(ws, row, fill=None, font_name=BODY_FONT, size=13):
    for col in range(COL['num'], COL['notes'] + 1):
        c = ws.cell(row=row, column=col)
        c.border = BORDER
        c.alignment = CENTER
        c.font = Font(name=font_name, size=size, bold=True)
        if fill:
            c.fill = PatternFill('solid', fgColor=fill)


def _write_header_labels(ws, row, fill):
    for col in range(COL['num'], COL['notes'] + 1):
        ws.cell(row=row, column=col, value=HEADERS[col - 1])
    _style_row(ws, row, fill=fill)


def _write_block(ws, row, block, first_num_literal, prev_num_row):
    """Write one group/committee block starting at `row`; return next free row."""
    girls = [p for p in block['people'] if p['gender'] == FEMALE]
    boys  = [p for p in block['people'] if p['gender'] == MALE]
    people_rows = max(1, len(girls) + len(boys))
    start, end = row, row + people_rows - 1

    L = get_column_letter
    fee = block['fee']
    # per-person rows: girls first (col E), then boys (col G)
    seq = [('girl', g) for g in girls] + [('boy', b) for b in boys]
    for i, (kind, person) in enumerate(seq):
        r = start + i
        if kind == 'girl':
            ws.cell(row=r, column=COL['girl'], value=person['name'])
        else:
            ws.cell(row=r, column=COL['boy'], value=person['name'])
        ws.cell(row=r, column=COL['gcount'], value=f'=IF(E{r}="",0,1)')
        ws.cell(row=r, column=COL['bcount'], value=f'=IF(G{r}="",0,1)')
        ws.cell(row=r, column=COL['fee'],
                value=(f'=(F{r}+H{r})*{fee}' if fee else 0))
        if person['notes']:
            ws.cell(row=r, column=COL['notes'], value=person['notes'])

    # block header cells (col B/C/D) on the first row
    if first_num_literal is not None:
        ws.cell(row=start, column=COL['num'], value=first_num_literal)
    elif prev_num_row:
        ws.cell(row=start, column=COL['num'],
                value=f'=MAX($B$6:B{prev_num_row})+1')
    ws.cell(row=start, column=COL['group'], value=block['title'])
    ws.cell(row=start, column=COL['total'],
            value=f'=SUM(F{start}:F{end},H{start}:H{end})')

    for r in range(start, end + 1):
        _style_row(ws, r)
    return end + 1


def _sum_row(ws, row, title, r0, r1):
    L = get_column_letter
    ws.cell(row=row, column=COL['group'], value=title)
    ws.cell(row=row, column=COL['total'], value=f'=SUM(D{r0}:D{r1})')
    ws.cell(row=row, column=COL['gcount'], value=f'=SUM(F{r0}:F{r1})')
    ws.cell(row=row, column=COL['bcount'], value=f'=SUM(H{r0}:H{r1})')
    ws.cell(row=row, column=COL['fee'],   value=f'=SUM(I{r0}:I{r1})')
    _style_row(ws, row)


def _place_image(ws, path, col0, row0, target_h=100, col_off=0, row_off=0):
    """Anchor an image at (col0,row0) — both 0-indexed — sized to `target_h` px
    tall with its NATIVE aspect ratio preserved (no stretching). A OneCellAnchor
    with an explicit extent is used so the display size is honoured (a plain
    string anchor makes openpyxl fall back to native pixels)."""
    if not (XLImage and path):
        return
    try:
        from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
        from openpyxl.drawing.xdr import XDRPositiveSize2D
        from openpyxl.utils.units import pixels_to_EMU
        img = XLImage(path)
        ratio = (img.width / img.height) if img.height else 1.0
        w_px = int(round(target_h * ratio))
        marker = AnchorMarker(col=col0, colOff=pixels_to_EMU(col_off),
                              row=row0, rowOff=pixels_to_EMU(row_off))
        img.anchor = OneCellAnchor(
            _from=marker,
            ext=XDRPositiveSize2D(pixels_to_EMU(w_px), pixels_to_EMU(target_h)))
        ws.add_image(img)
    except Exception:
        pass


def write_workbook(model, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Actual'
    ws.sheet_view.rightToLeft = True

    # print setup — fit all columns onto one page width (portrait A4), like the original
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = 9          # A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0        # unlimited pages tall
    ws.page_margins.left = ws.page_margins.right = 0
    ws.page_margins.top = 0.25
    ws.page_margins.bottom = 0.14

    for col, w in COL_WIDTHS.items():
        ws.column_dimensions[col].width = w

    # header rows 1-5
    ws.row_dimensions[1].height = 6.8
    _header_banner(ws, model['motto_text'], model['title'], model['location'])
    _write_header_labels(ws, 5, BLUE_FILL)

    # both logos: square source art, kept square, vertically centred in the banner
    # with a small inset from the sheet edges so they don't touch the outer border.
    _place_image(ws, model['jec_logo'],   col0=2, row0=0, target_h=86, row_off=15, col_off=10)   # right (C)
    _place_image(ws, model['motto_logo'], col0=9, row0=0, target_h=86, row_off=15, col_off=18)   # left  (J)

    # participants
    row = 6
    part_start = row
    first = True
    prev_num_row = None
    for block in model['participants']:
        used_from = row
        row = _write_block(ws, row, block,
                           first_num_literal=1 if first else None,
                           prev_num_row=prev_num_row)
        prev_num_row = row - 1        # last person row of this block
        first = False
        _separator_row(ws, row)       # coloured bar after the block
        row += 1
    part_end = row - 1
    _sum_row(ws, row, 'المجاميع', part_start, part_end)
    part_tot_row = row
    row += 2

    # committees section
    _write_header_labels(ws, row, GREEN_FILL)
    ws.cell(row=row, column=COL['group'], value='الأمانة واللجان')
    comm_start = row + 1
    row += 1
    for block in model['committees']:
        row = _write_block(ws, row, block, first_num_literal='{', prev_num_row=None)
        _separator_row(ws, row)
        row += 1
    comm_end = row - 1
    _sum_row(ws, row, 'مجاميع الأمانة واللجان', comm_start, comm_end)
    comm_tot_row = row
    row += 1

    # grand total
    L = get_column_letter
    ws.cell(row=row, column=COL['group'], value='المجاميع للمخيم')
    for key, col in (('total', 'D'), ('gcount', 'F'), ('bcount', 'H'), ('fee', 'I')):
        ws.cell(row=row, column=COL[key],
                value=f'={col}{part_tot_row}+{col}{comm_tot_row}')
    _style_row(ws, row)

    wb.save(out_path)


# ── event listing / selection ─────────────────────────────────────────────────
def _event_line(evt):
    reg = evt.get('registration', {}) or {}
    return '%s  %-8s %-16s  مشاركون=%d  مسؤولون=%d  أمانة=%d  ضيوف=%d' % (
        evt.get('id'), _txt(evt.get('event_type')), _txt(evt.get('title_label')),
        len(reg.get('members', [])), len(reg.get('supervisors', [])),
        len(reg.get('gs_committee', [])), len(reg.get('guests', [])))


def main():
    ap = argparse.ArgumentParser(description='Generate the camp participants roster xlsx for an event.')
    ap.add_argument('--event', help='Event id, e.g. EVT000002.')
    ap.add_argument('--out', help='Output .xlsx path. Default: tools/<event> roster.xlsx')
    ap.add_argument('--fee', type=float, default=None,
                    help='Participant subscription (per person). Default: event field or 25.')
    ap.add_argument('--resolutions', help='JSON file with pre-made gender/hull choices (non-interactive).')
    ap.add_argument('--no-open', action='store_true', help='Do not open the file when done.')
    ap.add_argument('--list', action='store_true', help='List events and exit.')
    args = ap.parse_args()

    _load_state_readonly()
    data = RE._load_events()
    events = data.get('events', [])

    if args.list or not args.event:
        for e in events:
            print(_event_line(e))
        if not args.event:
            print('\nPass --event <id> to generate.')
        return

    evt = next((e for e in events if e.get('id') == args.event), None)
    if not evt:
        print('No such event: %s' % args.event)
        sys.exit(1)

    fee = args.fee
    if fee is None:
        fee = evt.get('participant_fee')
    if fee is None:
        fee = 25
    fee = int(fee) if float(fee).is_integer() else fee

    out = args.out or os.path.join(_THIS_DIR, '%s roster.xlsx' % args.event)
    out = os.path.abspath(out)
    sidecar = args.resolutions or (out + '.resolutions.json')

    saved = {}
    if os.path.exists(sidecar):
        try:
            with open(sidecar, encoding='utf-8') as f:
                saved = json.load(f)
        except Exception:
            saved = {}

    interactive = (args.resolutions is None) and sys.stdin.isatty()
    resolver = Resolver(saved, interactive)

    model = build_model(evt, fee, resolver)

    if resolver.pending:
        print('\nCannot generate — unresolved choices (run interactively or supply --resolutions):')
        for line in resolver.pending:
            print('   • ' + line)
        print('\nresolutions file format: {"gender": {"<reg_id>": "ذكر|أنثى"}, '
              '"hull": {"<reg_id>": "<hull name>"}}')
        sys.exit(2)

    write_workbook(model, out)
    if args.resolutions is None:
        with open(sidecar, 'w', encoding='utf-8') as f:
            json.dump(resolver.saved, f, ensure_ascii=False, indent=1)

    n_part = sum(len(b['people']) for b in model['participants'])
    n_comm = sum(len(b['people']) for b in model['committees'])
    print('Wrote: %s' % out)
    print('  الشبيبات: %d  ·  مشاركون: %d' % (len(model['participants']), n_part))
    print('  اللجان:  %d  ·  أمانة/ضيوف: %d' % (len(model['committees']), n_comm))
    if not interactive and args.resolutions is None:
        print('  (non-interactive: all genders/hulls were already known or cached)')

    if not args.no_open:
        try:
            webbrowser.open('file:///' + out.replace('\\', '/'))
        except Exception:
            pass


if __name__ == '__main__':
    main()
